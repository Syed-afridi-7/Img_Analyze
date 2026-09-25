"""
Automated 15-Slide Presentation Deck Generator
Project: EXIF Metadata Extractor & Privacy Inspector (Img_Analyze)
Format: 16:9 Widescreen, Dark Cyberpunk Theme (Matching Sample PPT)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Initialize 16:9 Widescreen Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

# Cyberpunk Palette Definitions
COLOR_BG = RGBColor(11, 15, 23)         # #0B0F17
COLOR_CARD = RGBColor(20, 27, 39)       # #141B27
COLOR_CARD_BORDER = RGBColor(30, 41, 59)# #1E293B
COLOR_ACCENT = RGBColor(56, 189, 248)   # #38BDF8 (Cyan)
COLOR_TEXT = RGBColor(241, 245, 249)    # #F1F5F9 (White)
COLOR_MUTED = RGBColor(148, 163, 184)   # #94A3B8 (Gray)
COLOR_DANGER = RGBColor(248, 113, 113)  # #F87171 (Red)
COLOR_SUCCESS = RGBColor(52, 211, 153)  # #34D399 (Green)
COLOR_WARNING = RGBColor(251, 191, 36)  # #FBBF24 (Yellow)

def apply_background(slide):
    """Draw dark full-bleed background."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, subtitle_text, slide_num):
    """Render consistent header with slide pill."""
    # Title & Subtitle Box
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10.5), Inches(1.1))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT
    
    p2 = tf.add_paragraph()
    p2.text = subtitle_text.upper()
    p2.font.size = Pt(10)
    p2.font.color.rgb = COLOR_ACCENT
    p2.font.bold = True
    
    # Slide Number Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.8), Inches(0.45), Inches(0.75), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = COLOR_CARD
    badge.line.color.rgb = COLOR_ACCENT
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = f"{slide_num:02d}"
    p_b.font.size = Pt(12)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_ACCENT
    p_b.alignment = PP_ALIGN.CENTER

def add_footer(slide, category="EXIF Metadata Extractor & Privacy Inspector"):
    """Render standard footer."""
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.3))
    tf = txBox.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"SONA COLLEGE OF ARTS AND SCIENCE  •  BCA FINAL-YEAR PROJECT  •  {category.upper()}"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(100, 116, 139)

def create_card(slide, left, top, width, height, title, items, border_color=COLOR_CARD_BORDER, badge_text=None, badge_color=COLOR_ACCENT):
    """Reusable card component."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)
    
    # Card Header
    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT
    
    # Content Items
    for item in items:
        p = tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_TEXT
        p.space_before = Pt(5)

# ==========================================================
# SLIDE 1: TITLE SLIDE
# ==========================================================
s1 = prs.slides.add_slide(blank_slide_layout)
apply_background(s1)

# Pill banner
banner = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.3), Inches(1.0), Inches(4.7), Inches(0.45))
banner.fill.solid()
banner.fill.fore_color.rgb = COLOR_CARD
banner.line.color.rgb = COLOR_ACCENT
tf = banner.text_frame
p = tf.paragraphs[0]
p.text = "FINAL-YEAR PROJECT  •  CYBER SECURITY & FORENSICS"
p.font.size = Pt(10)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT
p.alignment = PP_ALIGN.CENTER

# Main Title & Subtitle
tx = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(2.2))
tf = tx.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "EXIF Metadata Extractor &\nPrivacy Inspector"
p.font.size = Pt(36)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT
p.alignment = PP_ALIGN.CENTER

p2 = tf.add_paragraph()
p2.text = "OSINT Image Telemetry Audit, Forensic Integrity Verification & In-Memory Privacy Sanitization Engine"
p2.font.size = Pt(14)
p2.font.color.rgb = COLOR_MUTED
p2.space_before = Pt(10)
p2.alignment = PP_ALIGN.CENTER

# Metadata Columns
meta_data = [
    ("SUBMITTED BY", "Student Name\nReg. No: C23UG206CAPxxx"),
    ("COURSE & DEPT", "Bachelor of Computer Applications\nDept. of Computer Applications"),
    ("INSTITUTION", "Sona College of Arts and Science\nPeriyar University, Salem"),
    ("ACADEMIC YEAR", "2025 – 2026\nGuide: Project Supervisor")
]
for idx, (head, val) in enumerate(meta_data):
    box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0 + idx * 2.9), Inches(4.5), Inches(2.6), Inches(1.4))
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_CARD
    box.line.color.rgb = COLOR_CARD_BORDER
    tf = box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = head
    p1.font.size = Pt(9)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_ACCENT
    p2 = tf.add_paragraph()
    p2.text = val
    p2.font.size = Pt(10)
    p2.font.color.rgb = COLOR_TEXT
    p2.space_before = Pt(4)

add_footer(s1, "Title Presentation")

# ==========================================================
# SLIDE 2: INTRODUCTION
# ==========================================================
s2 = prs.slides.add_slide(blank_slide_layout)
apply_background(s2)
add_header(s2, "Introduction: Digital Image Footprints", "The Context, Need & Project Purpose", 2)

create_card(s2, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.0), "THE CONTEXT", [
    "Modern smartphones and digital cameras record hidden metadata (EXIF) with every snapshot.",
    "Data silently travels with original image attachments when shared via email, cloud drives, or chat apps.",
    "Exposes exact capture environments, physical locations, and hardware signatures without user awareness."
])
create_card(s2, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.0), "THE PRIVACY NEED", [
    "Location privacy: Preserving home, school, and routine workplace coordinates from leaks.",
    "Forensic authenticity: Verifying whether digital media has been modified or re-compressed.",
    "Active mitigation: Providing users with instant, zero-trust tools to sanitize files before publishing."
])
create_card(s2, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.0), "PROJECT PURPOSE", [
    "A hybrid OSINT auditing and digital forensics privacy engine built with Python & Streamlit.",
    "Deeply parses EXIF, calculates cryptographic hashes, and maps GPS coordinates.",
    "Includes a zero-disk in-memory scrubber to strip metadata and export clean, safe photos."
])
add_footer(s2, "Introduction Phase")

# ==========================================================
# SLIDE 3: PROBLEM STATEMENT
# ==========================================================
s3 = prs.slides.add_slide(blank_slide_layout)
apply_background(s3)
add_header(s3, "Problem Statement: Metadata Vulnerabilities", "Unconscious Footprints & Cyber Risks", 3)

create_card(s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.4), "📍 Geolocation Exposure", [
    "Embedded GPS coordinates pinpoint individuals down to meters.",
    "Leads to physical stalking, burglary profiling, and residential tracing.",
    "Exposes children's school locations and personal daily travel routes."
], border_color=COLOR_DANGER)

create_card(s3, Inches(6.8), Inches(1.6), Inches(5.6), Inches(2.4), "📷 Hardware Profiling & Serials", [
    "Camera manufacturer, exact model, and lens specifications.",
    "Hardware serial numbers permanently attribute photos to owners.",
    "Firmware revisions expose unpatched software vulnerabilities."
])

create_card(s3, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.4), "⏰ Temporal Routine Tracking", [
    "High-precision timestamps record time down to the exact second.",
    "Aggregated photo collections reveal daily working hours and habits.",
    "Provides cyber-adversaries with chronological intelligence."
])

create_card(s3, Inches(6.8), Inches(4.3), Inches(5.6), Inches(2.4), "☁️ Cloud Scrubber Privacy Trap", [
    "Existing online metadata removers require uploading files to 3rd-party servers.",
    "Users must trade cloud storage privacy to remove local metadata.",
    "Creates secondary data leakage vectors and unauthorized retention."
], border_color=COLOR_DANGER)
add_footer(s3, "Problem Analysis")

# ==========================================================
# SLIDE 4: OBJECTIVES
# ==========================================================
s4 = prs.slides.add_slide(blank_slide_layout)
apply_background(s4)
add_header(s4, "Project Objectives & Technical Scope", "Measurable Engineering Deliverables", 4)

objectives = [
    ("1. Deep Forensic Extraction", "Extract and decode raw EXIF IFD tags into human-readable photography telemetry."),
    ("2. Cryptographic Integrity", "Compute in-memory MD5, SHA-1, and SHA-256 hashes to ensure chain of custody."),
    ("3. Geospatial Mapping", "Convert sexagesimal DMS coordinates to decimal degrees and render interactive maps."),
    ("4. AI Prompt Discovery", "Inspect PNG text chunks (tEXt, iTXt) to uncover Stable Diffusion / Midjourney prompts."),
    ("5. In-Memory Sanitization", "Strip all EXIF and GPS tags in-memory while preserving orientation via exif_transpose."),
    ("6. Multi-Format Reporting", "Generate multi-page PDF forensic audit reports (fpdf2), JSON payloads, and CSV matrices.")
]

for idx, (title, desc) in enumerate(objectives):
    row = idx // 2
    col = idx % 2
    create_card(s4, Inches(0.8 + col * 6.0), Inches(1.6 + row * 1.7), Inches(5.6), Inches(1.5), title, [desc])
add_footer(s4, "Project Objectives")

# ==========================================================
# SLIDE 5: EXISTING SYSTEM VS PROPOSED
# ==========================================================
s5 = prs.slides.add_slide(blank_slide_layout)
apply_background(s5)
add_header(s5, "Existing Solutions vs. Proposed Img_Analyze", "Comparative Architecture & Capabilities", 5)

create_card(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), "EXISTING SYSTEMS & LIMITATIONS", [
    "ExifTool / Command-Line: Steep learning curve for non-technical users; lacks visual maps and automated threat scoring.",
    "Online Web Uploaders: Requires uploading private media to remote 3rd-party cloud servers, worsening data exposure.",
    "Default OS Viewers: Display minimal data; no cryptographic hashing, flash bitmask parsing, or AI prompt detection.",
    "Destructive Strippers: Often delete orientation tags blindly, causing portrait photos to rotate or flip upside down upon sharing."
], border_color=COLOR_DANGER)

create_card(s5, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.1), "PROPOSED IMG_ANALYZE SYSTEM", [
    "100% In-Memory Processing: Air-gapped local execution via io.BytesIO. Zero disk footprint and zero telemetry.",
    "Dual-Interface Flexibility: Modern Streamlit Web Dashboard + lightweight CLI for automated terminal scripting.",
    "Multidimensional Forensics: Hashes, camera telemetry, GPS coordinates, 6-color palette, and AI generation parameters.",
    "Lossless Orientation Scrubber: Normalizes pixel transposition before EXIF strip to guarantee upright, clean sharing."
], border_color=COLOR_SUCCESS)
add_footer(s5, "System Study & Comparison")

# ==========================================================
# SLIDE 6: PROPOSED SYSTEM HIGHLIGHTS
# ==========================================================
s6 = prs.slides.add_slide(blank_slide_layout)
apply_background(s6)
add_header(s6, "Proposed System: Core Capabilities", "Forensic Auditing & Sanitization Suite", 6)

features = [
    ("🛡️ Zero-Disk Footprint", ["All files processed in RAM using io.BytesIO.", "Zero temp files or disk artifacts left on machine."]),
    ("📍 Interactive Geocoding", ["DMS to decimal conversion with hemisphere handling.", "Direct PyDeck maps + OpenStreetMap iframe layers."]),
    ("🎨 Palette Quantization", ["MEDIANCUT algorithm surfaces top 6 dominant colors.", "Perceptual luminance & RMS contrast classification."]),
    ("🤖 AI Generation Tracing", ["Inspects PNG chunks for Stable Diffusion & Midjourney.", "Extracts positive prompts, negative prompts, and seeds."]),
    ("🔄 Lossless Transpose", ["Auto-corrects photo rotation using exif_transpose.", "Prevents upside-down images after metadata strip."]),
    ("📑 Forensic PDF Export", ["Multi-page formatted PDF report using fpdf2.", "Structured JSON & batch CSV matrix download."])
]
for idx, (title, items) in enumerate(features):
    row = idx // 3
    col = idx % 3
    create_card(s6, Inches(0.8 + col * 4.0), Inches(1.6 + row * 2.6), Inches(3.7), Inches(2.3), title, items)
add_footer(s6, "Proposed System Architecture")

# ==========================================================
# SLIDE 7: SYSTEM ARCHITECTURE
# ==========================================================
s7 = prs.slides.add_slide(blank_slide_layout)
apply_background(s7)
add_header(s7, "System Architecture & Processing Pipeline", "Ingestion ➔ Extraction ➔ Sanitization", 7)

arch_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.1))
arch_box.fill.solid()
arch_box.fill.fore_color.rgb = COLOR_CARD
arch_box.line.color.rgb = COLOR_CARD_BORDER
tf = arch_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "DATA FLOW & PIPELINE TOPOLOGY"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

steps = [
    ("1. Ingestion Layer", "Supports drag-and-drop file stream directly into memory (JPEG, PNG, WebP, TIFF, BMP, GIF)."),
    ("2. Cryptographic Engine", "Computes MD5, SHA-1, and SHA-256 hashes in-memory for forensic chain-of-custody verification."),
    ("3. Telemetry Parser", "Extracts raw IFD blocks; decodes Aperture (f-stop), Shutter (fractions), ISO, and Flash bitmasks."),
    ("4. Geolocation Engine", "Converts DMS tuples to signed decimal coordinates and renders OpenStreetMap / PyDeck layers."),
    ("5. Sanitization Engine", "Transposes pixels to upright orientation, strips EXIF/APP1 blocks, and outputs a clean file.")
]
for title, desc in steps:
    p = tf.add_paragraph()
    p.text = f"▶  {title}:  {desc}"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT
    p.space_before = Pt(12)
add_footer(s7, "Architectural Design")

# ==========================================================
# SLIDE 8: FORENSIC WORKFLOW & CRYPTOGRAPHY
# ==========================================================
s8 = prs.slides.add_slide(blank_slide_layout)
apply_background(s8)
add_header(s8, "Cryptographic Hashes & Telemetry Decoding", "Forensic Integrity & Bitmask Analysis", 8)

create_card(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), "CRYPTOGRAPHIC CHAIN OF CUSTODY", [
    "Calculates MD5, SHA-1, and SHA-256 checksums in-memory directly from raw stream.",
    "Establishes tamper-evident baseline before performing any metadata extraction.",
    "Ensures that sanitized outputs can be mathematically proven distinct from originals.",
    "Supports digital forensics law-enforcement evidentiary standards."
])
create_card(s8, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.1), "FLASH BITMASK & OPTICS DECODING", [
    "Raw EXIF Flash tags store complex operational states inside single integer bitmasks.",
    "Bit 0: Flash fired / did not fire.",
    "Bits 1-2: Strobe return light detection status.",
    "Bits 3-4: Compulsory flash mode / suppressive mode.",
    "Bit 6: Red-eye reduction mode enabled.",
    "Translates raw shutter timings (e.g., 0.004) into human photographic fractions (1/250 s)."
])
add_footer(s8, "Forensic Cryptography & Optics")

# ==========================================================
# SLIDE 9: GEOLOCATION TRANSFORMATION
# ==========================================================
s9 = prs.slides.add_slide(blank_slide_layout)
apply_background(s9)
add_header(s9, "Geolocation & Coordinate Mathematics", "Sexagesimal DMS to Decimal Degree Conversion", 9)

create_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), "MATHEMATICAL FORMULA", [
    "EXIF GPS tags store coordinates as rational degree, minute, and second tuples.",
    "Conversion Algorithm:",
    "Decimal = Degrees + (Minutes / 60.0) + (Seconds / 3600.0)",
    "Hemispheric Negation:",
    "If GPSLatitudeRef == 'S': Decimal = -Decimal",
    "If GPSLongitudeRef == 'W': Decimal = -Decimal",
    "Altitude Calculation: Resolves altitude in meters referencing sea level flags."
])
create_card(s9, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.1), "MAPPING & RECONNAISSANCE LAYER", [
    "Native PyDeck Map: Embeds pinpointed vector markers directly inside the UI.",
    "OpenStreetMap Iframe: Loads street-level cartographic tiles without API keys.",
    "One-Click Navigators: Direct hyperlinks to Google Maps, Apple Maps, and OSM.",
    "High Risk Warning: Prominently warns users if photos expose home coordinates."
], border_color=COLOR_DANGER)
add_footer(s9, "Geolocation Engine")

# ==========================================================
# SLIDE 10: SYSTEM MODULES BREAKDOWN
# ==========================================================
s10 = prs.slides.add_slide(blank_slide_layout)
apply_background(s10)
add_header(s10, "Modular Subsystem Breakdown", "Core Codebase Components", 10)

modules = [
    ("extractor.py", "Core extraction engine: parses IFDs, color spaces, bit depth, and generates hashes."),
    ("formatter.py", "Telemetry translator: decodes shutter fractions, apertures, and bitmask states."),
    ("gps.py", "Geospatial engine: converts DMS to signed decimals and constructs navigation URLs."),
    ("batch.py", "Batch coordinator: multi-image matrices, aggregate stats, and multi-point map plots."),
    ("pdf_export.py", "Document builder: generates formatted multi-page digital forensic PDF audit reports."),
    ("app.py / cli.py", "Dual interfaces: Interactive Streamlit web dashboard + headless CLI script.")
]
for idx, (title, desc) in enumerate(modules):
    row = idx // 2
    col = idx % 2
    create_card(s10, Inches(0.8 + col * 6.0), Inches(1.6 + row * 1.7), Inches(5.6), Inches(1.5), title, [desc])
add_footer(s10, "Module Specifications")

# ==========================================================
# SLIDE 11: TECHNOLOGIES & TOOLS USED
# ==========================================================
s11 = prs.slides.add_slide(blank_slide_layout)
apply_background(s11)
add_header(s11, "Technologies & Environmental Stack", "Programming Stack & Architecture", 11)

tech_stack = [
    ("Python 3.10+", ["Core programming language.", "Robust in-memory byte stream handling."]),
    ("Pillow (PIL 10.0+)", ["Advanced image processing.", "ExifTags, ImageOps & color quantization."]),
    ("Streamlit 1.30+", ["Interactive web dashboard.", "Native tabs, file dropzone & PyDeck maps."]),
    ("fpdf2 2.8+", ["Multi-page PDF generation engine.", "Constructs formal forensic audit reports."]),
    ("Pandas & NumPy", ["Dataframe manipulation.", "Generates side-by-side batch comparison matrices."]),
    ("Standard Hashlib", ["In-memory cryptographic digests.", "Calculates MD5, SHA-1, and SHA-256."])
]
for idx, (title, items) in enumerate(tech_stack):
    row = idx // 3
    col = idx % 3
    create_card(s11, Inches(0.8 + col * 4.0), Inches(1.6 + row * 2.6), Inches(3.7), Inches(2.3), title, items)
add_footer(s11, "Technology Specifications")

# ==========================================================
# SLIDE 12: VISUAL ANALYTICS & AI PROMPTS
# ==========================================================
s12 = prs.slides.add_slide(blank_slide_layout)
apply_background(s12)
add_header(s12, "Visual Analytics & AI Prompt Tracing", "Computer Vision & Generative Media Forensics", 12)

create_card(s12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), "COLOR PALETTE & LIGHTING ANALYSIS", [
    "Dominant Colors: Uses Pillow's MEDIANCUT quantization to isolate the top 6 dominant colors.",
    "Outputs Hex color codes, RGB channels, and precise surface coverage percentages.",
    "Perceptual Luminance: Computes brightness via Rec. 601 formula (Y = 0.299R + 0.587G + 0.114B).",
    "RMS Contrast: Classifies lighting into High-Key, Low-Key, or Balanced exposure."
])
create_card(s12, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.1), "AI GENERATION CHUNK INSPECTION", [
    "AI generators (Stable Diffusion, Midjourney, ComfyUI) embed metadata inside PNG text chunks.",
    "Engine parses tEXt, zTXt, and iTXt ancillary chunks.",
    "Extracts positive generation prompts, negative prompts, seed values, CFG scale, and steps.",
    "Provides immediate forensic attribution for synthetic/AI-generated media."
], border_color=COLOR_WARNING)
add_footer(s12, "Visual & AI Forensics")

# ==========================================================
# SLIDE 13: TESTING & EXPERIMENTAL RESULTS
# ==========================================================
s13 = prs.slides.add_slide(blank_slide_layout)
apply_background(s13)
add_header(s13, "System Testing & Quality Assurance", "Automated Test Matrix (48 Test Cases)", 13)

create_card(s13, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), "AUTOMATED TEST SUITE COVERAGE", [
    "test_image_details.py: Validates dimensions, bit depth, color modes, and hash integrity across 6 formats.",
    "test_batch.py: Tests multi-file aggregation, risk categorization, and comparative matrix calculations.",
    "test_pdf_export.py: Validates multi-page PDF generation, table formatting, and Unicode stability.",
    "test_streamlit_app.py: Headless Streamlit AppTest validating UI reactivity, widgets, and tab switching."
])
create_card(s13, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.1), "EMPIRICAL EVALUATION RESULTS", [
    "Extraction Accuracy: 100% EXIF tag retention with zero unhandled exceptions.",
    "Scrubbing Verification: Confirmed 0 bytes of residual metadata post-sanitization.",
    "Orientation Invariance: Flawlessly preserved vertical portrait shots using exif_transpose.",
    "Network Privacy: Zero telemetry verified (gatherUsageStats = false)."
], border_color=COLOR_SUCCESS)
add_footer(s13, "Testing & Verification")

# ==========================================================
# SLIDE 14: CONCLUSION & FUTURE SCOPE
# ==========================================================
s14 = prs.slides.add_slide(blank_slide_layout)
apply_background(s14)
add_header(s14, "Conclusion & Future Enhancements", "Research Summary & Technical Roadmap", 14)

create_card(s14, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), "PROJECT CONCLUSION", [
    "Successfully built an end-to-end OSINT metadata auditing and in-memory sanitization suite.",
    "Combines cryptographic chain-of-custody tracking with interactive geocoding and visual analytics.",
    "Eliminates reliance on insecure third-party cloud uploaders via 100% local in-memory execution.",
    "Provides users and forensic auditors with complete digital privacy autonomy."
])
create_card(s14, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.1), "FUTURE ENHANCEMENTS", [
    "Computer Vision Redaction: Integrate local YOLO/OpenCV models to blur visible faces and license plates.",
    "Camera RAW Support: Expand native decoding to proprietary formats (.CR2, .NEF, .ARW, .DNG).",
    "Video Container Forensics: Extend atom parsers to MP4, MOV, and MKV video container telemetry.",
    "Automated CLI Daemon: Background folder watcher for instant automated scrubbing."
], border_color=COLOR_ACCENT)
add_footer(s14, "Concluding Summary")

# ==========================================================
# SLIDE 15: APPLICATION SCREENSHOTS (EXACT SCREENSHOTS GRID)
# ==========================================================
s15 = prs.slides.add_slide(blank_slide_layout)
apply_background(s15)
add_header(s15, "Application Screenshots: Live System Demonstration", "Verified Execution Views & Audit Workflows", 15)

# Map user's uploaded screenshots
screenshots = [
    ("Screenshot 2026-09-25 103036.jpg", "1. Home Dashboard & Dropzone"),
    ("Screenshot 2026-09-25 103105.jpg", "2. In-Memory Image Selection"),
    ("Screenshot 2026-09-25 103141.jpg", "3. High Risk Alert (GPS Leak)"),
    ("Screenshot 2026-09-25 103125.jpg", "4. Low Risk Sanitized Verification"),
    ("Screenshot 2026-09-25 103227.png", "5. File Specs & Hashing"),
    ("Screenshot 2026-09-25 103235.jpg", "6. Camera Telemetry View")
]

# Grid of 6 screenshots (3 cols x 2 rows)
card_w = Inches(3.7)
card_h = Inches(2.3)

for idx, (img_filename, caption) in enumerate(screenshots):
    col = idx % 3
    row = idx // 3
    left = Inches(0.8 + col * 4.0)
    top = Inches(1.6 + row * 2.5)
    
    # Outer frame card
    box = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_CARD
    box.line.color.rgb = COLOR_CARD_BORDER
    
    # Check if screenshot image exists in current folder
    if os.path.exists(img_filename):
        try:
            s15.shapes.add_picture(img_filename, left + Inches(0.08), top + Inches(0.08), card_w - Inches(0.16), card_h - Inches(0.48))
        except Exception:
            pass
    
    # Caption at bottom of card
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = caption
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(140)

# Slide 15 Final Footer banner
footer_box = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.45))
footer_box.fill.solid()
footer_box.fill.fore_color.rgb = COLOR_CARD
footer_box.line.color.rgb = COLOR_ACCENT
tf = footer_box.text_frame
p = tf.paragraphs[0]
p.text = "THANK YOU  •  QUESTIONS & DISCUSSION  •  BCA FINAL YEAR PROJECT VIVA"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT
p.alignment = PP_ALIGN.CENTER

# Save presentation
output_path = "EXIF_Metadata_Extractor_15_Slides.pptx"
prs.save(output_path)
print(f"[SUCCESS] Successfully generated perfect 15-Slide Presentation: {output_path}")