"""
generate_bca_docx.py
Converts the official BCA Academic Mini Project Report markdown files in docs/bca_report/
into a publication-grade, beautifully formatted Microsoft Word document: final_pro.docx
Adheres strictly to the university report structure:
Sona College of Arts and Science (Affiliated to Periyar University, Salem)
Calibrated strictly to 75-80 pages.
"""

import os
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BCA_REPORT_DIR = r"D:\IMG_ANALYZE\docs\bca_report"
OUTPUT_DOCX = r"D:\IMG_ANALYZE\final_pro.docx"

# Academic Corporate Palette
NAVY = RGBColor(27, 54, 93)       # #1B365D - Deep University Navy
SLATE = RGBColor(44, 94, 138)     # #2C5E8A - Secondary Steel Slate
CHARCOAL = RGBColor(45, 45, 45)   # #2D2D2D - Dark Readable Body Text
MUTED = RGBColor(100, 110, 120)   # #646E78 - Captions, Headers, Footers
CODE_GRAY = RGBColor(30, 30, 30)  # #1E1E1E - Code Block Monospace

HEX_NAVY = "1B365D"
HEX_LIGHT_BG = "F4F6F9"
HEX_CODE_BG = "F6F8FA"

# Typography & Sizing baseline (can be tuned dynamically)
MARGIN_IN = 0.55
BODY_PT = 8.3
LINE_SP = 0.98
SPACE_AFTER_PT = 0.8
TABLE_PT = 7.3
CODE_PT = 6.5


def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=10, bottom=10, left=30, right=30):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)


def set_table_borders(table, color="D0D7DE"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="1B365D"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)


def add_formatted_runs(paragraph, text, is_code=False, font_size_pt=BODY_PT, code_font_size=CODE_PT):
    if is_code:
        run = paragraph.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(code_font_size)
        run.font.color.rgb = CODE_GRAY
        return

    tokens = re.split(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('***') and token.endswith('***') and len(token) >= 6:
            run = paragraph.add_run(token[3:-3])
            run.bold = True
            run.italic = True
            run.font.size = Pt(font_size_pt)
        elif token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.size = Pt(font_size_pt)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.size = Pt(font_size_pt)
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(font_size_pt - 0.5)
            run.font.color.rgb = RGBColor(160, 40, 60)
        else:
            run = paragraph.add_run(token)
            run.font.size = Pt(font_size_pt)


def add_page_number(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)


def add_total_pages(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> NUMPAGES </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)


def create_document(margin_in=MARGIN_IN, body_pt=BODY_PT, line_sp=LINE_SP, space_after_pt=SPACE_AFTER_PT):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(margin_in)
        section.bottom_margin = Inches(margin_in)
        section.left_margin = Inches(margin_in)
        section.right_margin = Inches(margin_in)

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("EXIF Metadata Extractor & Privacy Inspector | BCA Academic Mini Project Report")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(7.5)
        hrun.font.color.rgb = MUTED

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        frun1 = fp.add_run("Sona College of Arts and Science (Periyar University)  |  Page ")
        frun1.font.name = "Calibri"
        frun1.font.size = Pt(7.5)
        frun1.font.color.rgb = MUTED

        page_run = fp.add_run()
        page_run.font.name = "Calibri"
        page_run.font.size = Pt(7.5)
        page_run.font.color.rgb = NAVY
        page_run.bold = True
        add_page_number(page_run)

        frun2 = fp.add_run(" of ")
        frun2.font.name = "Calibri"
        frun2.font.size = Pt(7.5)
        frun2.font.color.rgb = MUTED

        total_run = fp.add_run()
        total_run.font.name = "Calibri"
        total_run.font.size = Pt(7.5)
        total_run.font.color.rgb = MUTED
        add_total_pages(total_run)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(body_pt)
    style_normal.font.color.rgb = CHARCOAL
    style_normal.paragraph_format.line_spacing = line_sp
    style_normal.paragraph_format.space_after = Pt(space_after_pt)

    return doc


def process_chapter_file(doc, filepath, is_first_file=False, body_pt=BODY_PT, table_pt=TABLE_PT, code_pt=CODE_PT, line_sp=LINE_SP, space_after_pt=SPACE_AFTER_PT):
    filename = os.path.basename(filepath)
    print(f"Processing: {filename}...")

    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # Major chapter boundary page break
    if not is_first_file and filename.startswith(('01_', '02_', '03_', '04_', '05_', '06_', '07_')):
        doc.add_page_break()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
        code_text = "".join(code_lines)
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_shading(cell, HEX_CODE_BG)
        set_cell_margins(cell, top=10, bottom=10, left=40, right=40)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.line_spacing = 0.95
        add_formatted_runs(cp, code_text, is_code=True, font_size_pt=body_pt, code_font_size=code_pt)

        code_lines = []
        in_code_block = False

    def flush_tbl():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return

        parsed_rows = []
        for tline in table_lines:
            raw = tline.strip()
            if not raw or not raw.startswith('|'):
                continue
            cells = [c.strip() for c in raw.split('|')[1:-1]]
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            parsed_rows.append(cells)

        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            for r in parsed_rows:
                while len(r) < num_cols:
                    r.append("")

            table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table)

            for row_idx, row_data in enumerate(parsed_rows):
                row = table.rows[row_idx]
                is_header = (row_idx == 0)
                for col_idx, cell_value in enumerate(row_data):
                    cell = row.cells[col_idx]
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    set_cell_margins(cell, top=10, bottom=10, left=30, right=30)

                    if is_header:
                        set_cell_shading(cell, HEX_NAVY)
                    elif row_idx % 2 == 1:
                        set_cell_shading(cell, HEX_LIGHT_BG)

                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 0.98
                    add_formatted_runs(p, cell_value, font_size_pt=table_pt)

                    if is_header:
                        for run in p.runs:
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.bold = True
                            run.font.size = Pt(table_pt)

        table_lines = []
        in_table = False

    for line in lines:
        stripped = line.strip()

        # Explicit markdown page break indicator (e.g. preliminary pages)
        if '<div style="page-break-after: always;"></div>' in stripped:
            if in_code_block: flush_code()
            if in_table: flush_tbl()
            doc.add_page_break()
            continue

        if stripped == '---':
            if in_code_block: flush_code()
            if in_table: flush_tbl()
            continue

        if stripped.startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table: flush_tbl()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(line)
            continue
        elif in_table:
            flush_tbl()

        if not stripped:
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[2:].strip())
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(body_pt + 5.0)
            run.font.color.rgb = NAVY
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[3:].strip())
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(body_pt + 3.0)
            run.font.color.rgb = NAVY
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(0.5)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:].strip())
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(body_pt + 1.5)
            run.font.color.rgb = SLATE
        elif stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(0.5)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[5:].strip())
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(body_pt + 0.5)
            run.font.color.rgb = CHARCOAL
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.right_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run_bar = p.add_run("┃ ")
            run_bar.bold = True
            run_bar.font.color.rgb = NAVY
            add_formatted_runs(p, stripped[2:].strip(), font_size_pt=body_pt - 0.5)
            for r in p.runs[1:]:
                r.italic = True
        elif stripped.startswith(("- ", "* ", "• ")):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0.5)
            p.paragraph_format.line_spacing = line_sp
            add_formatted_runs(p, stripped[2:].strip(), font_size_pt=body_pt)
        elif re.match(r'^\d+\.\s+', stripped):
            num_match = re.match(r'^\d+\.\s+', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0.5)
            p.paragraph_format.line_spacing = line_sp
            add_formatted_runs(p, stripped[num_match.end():].strip(), font_size_pt=body_pt)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(space_after_pt)
            p.paragraph_format.line_spacing = line_sp
            add_formatted_runs(p, stripped, font_size_pt=body_pt)

    if in_code_block: flush_code()
    if in_table: flush_tbl()


def build_docx(output_path=OUTPUT_DOCX, margin_in=MARGIN_IN, body_pt=BODY_PT, table_pt=TABLE_PT, code_pt=CODE_PT, line_sp=LINE_SP, space_after_pt=SPACE_AFTER_PT):
    print("="*70)
    print(f"Building BCA Project Report: {output_path}")
    print(f"Config: margin={margin_in}in, body={body_pt}pt, table={table_pt}pt, code={code_pt}pt, line_sp={line_sp}, space_after={space_after_pt}pt")
    print("="*70)

    doc = create_document(margin_in=margin_in, body_pt=body_pt, line_sp=line_sp, space_after_pt=space_after_pt)

    ordered_files = [
        "00_preliminary.md",
        "01_introduction.md",
        "02_system_study.md",
        "03_system_design_development.md",
        "04_testing_implementation.md",
        "05_conclusion_future_scope.md",
        "06_bibliography.md",
        "07_appendices.md"
    ]

    for idx, fname in enumerate(ordered_files):
        fpath = os.path.join(BCA_REPORT_DIR, fname)
        if os.path.exists(fpath):
            process_chapter_file(doc, fpath, is_first_file=(idx == 0), body_pt=body_pt, table_pt=table_pt, code_pt=code_pt, line_sp=line_sp, space_after_pt=space_after_pt)
        else:
            print(f"Warning: File {fname} not found!")

    doc.save(output_path)
    size_mb = os.path.getsize(output_path) / (1024 * 1024)
    print(f"Document saved to {output_path} ({size_mb:.2f} MB)")


def get_word_page_count(docx_path):
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(os.path.abspath(docx_path))
        pages = doc.ComputeStatistics(2) # 2 = wdStatisticPages
        doc.Close(False)
        word.Quit()
        return pages
    except Exception as e:
        print(f"Error computing Word statistics via COM: {e}")
        return None


if __name__ == "__main__":
    build_docx()
    pages = get_word_page_count(OUTPUT_DOCX)
    print(f"\n>>> OFFICIAL WORD COM RENDERED PAGE COUNT: {pages} PAGES <<<")
